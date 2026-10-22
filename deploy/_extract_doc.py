# -*- coding: utf-8 -*-
from docx import Document
from pathlib import Path

p = Path(r"C:\Users\Ibiza\Desktop\部署说明.docx")
d = Document(str(p))
out = Path(r"C:\Users\Ibiza\Desktop\project\挑战杯\deploy\_deploy_doc_extract.txt")
lines = []
for i, para in enumerate(d.paragraphs):
    style = para.style.name if para.style else ""
    lines.append(f"{i}\t{style}\t{para.text}")
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "paras", len(d.paragraphs))
